# -*- coding: utf-8 -*-
import os
import time
import cv2
import gradio as gr
import torch
import xml.etree.ElementTree as ET
from ultralytics import YOLO
import pygame
import openpyxl
from openpyxl.drawing.image import Image as ExcelImage
import base64

device = "cuda" if torch.cuda.is_available() else "cpu"

# 🛠️ 修正 1：移除 "project/" 前綴
model = YOLO("models/最新模型.pt")

pygame.mixer.init()

# 🛠️ 修正 2：移除 "project/" 前綴
LOG_IMG_DIR = "log_images"
os.makedirs(LOG_IMG_DIR, exist_ok=True)
EXCEL_EXPORT_DIR = "exports"
os.makedirs(EXCEL_EXPORT_DIR, exist_ok=True)

# ==========================================
# 📊 報表與介面輔助函數
# ==========================================
def generate_html_table(records, headers=["時間", "訊息內容", "截圖"]):
    html = "<table style='width:100%; border-collapse: collapse;'>"
    html += "<tr style='background-color: #f2f2f2;'>"
    for h in headers:
        html += f"<th style='border: 1px solid #ddd; padding: 8px; text-align: left;'>{h}</th>"
    html += "</tr>"
    
    for row in records:
        html += "<tr>"
        for i, cell in enumerate(row):
            if i == 2 and cell: # 截圖欄位
                if os.path.exists(cell):
                    with open(cell, "rb") as f:
                        b64_img = base64.b64encode(f.read()).decode('utf-8')
                    html += f"<td style='border: 1px solid #ddd; padding: 8px;'><img src='data:image/jpeg;base64,{b64_img}' width='150' style='border-radius:5px;'></td>"
                else:
                    html += f"<td style='border: 1px solid #ddd; padding: 8px;'>找不到圖片</td>"
            else:
                html += f"<td style='border: 1px solid #ddd; padding: 8px;'>{cell}</td>"
        html += "</tr>"
    html += "</table>"
    return html

def export_to_excel(records, filepath, headers=["時間", "訊息內容", "截圖"]):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 30

    for r_idx, row_data in enumerate(records, start=2):
        time_str, msg, img_path = row_data
        ws.cell(row=r_idx, column=1, value=time_str)
        ws.cell(row=r_idx, column=2, value=msg)
        
        if img_path and os.path.exists(img_path):
            img = ExcelImage(img_path)
            img.width = 160
            img.height = 120
            ws.row_dimensions[r_idx].height = 95 
            ws.add_image(img, f"C{r_idx}")
        else:
            ws.cell(row=r_idx, column=3, value="無圖片")
            ws.row_dimensions[r_idx].height = 20
            
    wb.save(filepath)

# ==========================================
# 🎨 UI 輔助繪圖函數
# ==========================================
def draw_warning_text(img, text="不完備"):
    font = cv2.FONT_HERSHEY_SIMPLEX
    scale = 1.0
    thickness = 3
    (text_width, text_height), baseline = cv2.getTextSize(text, font, scale, thickness)
    x, y = 50, 50
    pad = 10
    
    cv2.rectangle(img, 
                  (x - pad, y - text_height - pad), 
                  (x + text_width + pad, y + baseline + pad), 
                  (255, 255, 255), -1)
    cv2.putText(img, text, (x, y), font, scale, (0, 0, 255), thickness, cv2.LINE_AA)

def draw_cached_boxes(img, boxes_info):
    for box in boxes_info:
        cls_name, conf, x1, y1, x2, y2 = box
        if "no_" in cls_name.lower():
            color = (0, 0, 255)  
            txt_color = (255, 255, 255)
        else:
            color = (0, 255, 0)  
            txt_color = (0, 0, 0)

        cv2.rectangle(img, (int(x1), int(y1)), (int(x2), int(y2)), color, 2)
        label = f"{cls_name} {conf:.2f}"
        (tw, th), _ = cv2.getTextSize(label, cv2.FONT_HERSHEY_SIMPLEX, 0.5, 1)
        cv2.rectangle(img, (int(x1), int(y1) - th - 5), (int(x1) + tw, int(y1)), color, -1)
        cv2.putText(img, label, (int(x1), int(y1) - 3), cv2.FONT_HERSHEY_SIMPLEX, 0.5, txt_color, 1, cv2.LINE_AA)

# ==========================================
# 📁 資料儲存邏輯
# ==========================================
# 🛠️ 修正 3：預設儲存路徑改為 "data"
def save_voc_annotation(img_name, img_width, img_height, boxes_info, save_dir="data"):
    os.makedirs(save_dir, exist_ok=True)
    
    annotation = ET.Element("annotation")
    ET.SubElement(annotation, "folder").text = save_dir
    ET.SubElement(annotation, "filename").text = img_name
    
    size = ET.SubElement(annotation, "size")
    ET.SubElement(size, "width").text = str(img_width)
    ET.SubElement(size, "height").text = str(img_height)
    ET.SubElement(size, "depth").text = "3"
    
    for box in boxes_info:
        cls_name = box[0]
        xmin, ymin, xmax, ymax = box[2], box[3], box[4], box[5]
        
        obj = ET.SubElement(annotation, "object")
        ET.SubElement(obj, "name").text = cls_name
        ET.SubElement(obj, "pose").text = "Unspecified"
        ET.SubElement(obj, "truncated").text = "0"
        ET.SubElement(obj, "difficult").text = "0"
        
        bndbox = ET.SubElement(obj, "bndbox")
        ET.SubElement(bndbox, "xmin").text = str(int(xmin))
        ET.SubElement(bndbox, "ymin").text = str(int(ymin))
        ET.SubElement(bndbox, "xmax").text = str(int(xmax))
        ET.SubElement(bndbox, "ymax").text = str(int(ymax))
        
    tree = ET.ElementTree(annotation)
    xml_filename = os.path.join(save_dir, img_name.rsplit(".", 1)[0] + ".xml")
    tree.write(xml_filename, encoding="utf-8", xml_declaration=True)

# ==========================================
# 🚀 核心：本機端 OpenCV 攝影機迴圈
# ==========================================
def run_local_webcam(interval, confidence_level, alarm_enabled, dataOff_set, voice_class, voice_vol, history_logs):
    cap = cv2.VideoCapture(0)
    if not cap.isOpened():
        err_logs = history_logs.copy() if history_logs else []
        err_logs.append([time.strftime('%H:%M:%S'), "❌ 無法開啟本機攝影機", ""])
        yield generate_html_table(err_logs), None, err_logs
        return

    log_records = history_logs.copy() if history_logs else []
    webcam_counter = 0
    sin_token = 0
    is_alerting = False
    cached_boxes = [] 
    
    interval = int(interval) if interval and interval > 0 else 4
    conf_threshold = float(confidence_level) / 100.0 if confidence_level else 0.50

    # 🛠️ 修正 4：移除 "project/" 前綴
    audio_path = f"voice/{voice_class}.mp3"
    alarm_sound = None
    if os.path.exists(audio_path):
        alarm_sound = pygame.mixer.Sound(audio_path)
        alarm_sound.set_volume(voice_vol / 100.0)

    excel_filename = f"webcam_log_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    excel_filepath = os.path.join(EXCEL_EXPORT_DIR, excel_filename)

    def push_update():
        export_to_excel(log_records, excel_filepath)
        return generate_html_table(log_records), excel_filepath, log_records

    log_records.append([time.strftime('%H:%M:%S'), f"📷 攝影機已啟動 (門檻: {conf_threshold*100:.0f}%)", ""])
    yield push_update()
    print("💡 在影像視窗按下鍵盤上的 'q' 鍵即可關閉攝影機。")

    while cap.isOpened():
        ret, frame = cap.read()
        if not ret:
            break

        webcam_counter += 1
        output_frame = frame.copy()
        log_updated = False 

        if (webcam_counter - 1) % interval == 0:
            results = model(frame, conf=conf_threshold)
            detected_classes = [model.names[int(cls)].lower() for cls in results[0].boxes.cls]
            
            cached_boxes = []
            for box in results[0].boxes:
                cls_id = int(box.cls[0])
                cls_name = model.names[cls_id]
                conf = float(box.conf[0])
                xyxy = box.xyxy[0].tolist()
                cached_boxes.append((cls_name, conf, xyxy[0], xyxy[1], xyxy[2], xyxy[3]))

            if dataOff_set and len(results[0].boxes) > 0 and "person" in detected_classes:
                h, w = frame.shape[:2]
                timestamp = time.strftime("%Y%m%d_%H%M%S") + f"_{int(time.time()*1000)%1000}"
                img_filename = f"train_{timestamp}.jpg"
                
                # 🛠️ 修正 5：直接設定為 "data"
                save_dir = "data"
                os.makedirs(save_dir, exist_ok=True)
                img_path = os.path.join(save_dir, img_filename)
                
                cv2.imwrite(img_path, frame)
                save_voc_annotation(img_filename, w, h, cached_boxes, save_dir=save_dir)

            if "no_helmet" in detected_classes or "no_gloves" in detected_classes or "no_glove" in detected_classes:
                sin_token += 1
                if sin_token == 4 and not is_alerting:
                    is_alerting = True
                    
                    temp_frame = output_frame.copy()
                    draw_cached_boxes(temp_frame, cached_boxes)
                    img_name = f"alert_{time.strftime('%Y%m%d_%H%M%S')}.jpg"
                    img_path = os.path.join(LOG_IMG_DIR, img_name)
                    cv2.imwrite(img_path, temp_frame)
                    
                    log_records.append([time.strftime('%H:%M:%S'), "⚠️ 偵測到消防員未配戴完整裝備！", img_path])
                    log_updated = True
                    
                    if alarm_enabled and alarm_sound:
                        alarm_sound.play(-1)
            else:
                sin_token = 0 
                if is_alerting:
                    is_alerting = False
                    log_records.append([time.strftime('%H:%M:%S'), "✅ 裝備已恢復完整，警報解除。", ""])
                    log_updated = True
                    if alarm_sound:
                        alarm_sound.stop()

        draw_cached_boxes(output_frame, cached_boxes)

        if is_alerting and alarm_enabled:
            draw_warning_text(output_frame)

        cv2.imshow("Live Detection (Press 'q' to quit)", output_frame)

        if log_updated:
            yield push_update()

        if cv2.waitKey(1) & 0xFF == ord('q'):
            log_records.append([time.strftime('%H:%M:%S'), "⏹️ 攝影機已手動關閉。", ""])
            yield push_update()
            break

    if alarm_sound:
        alarm_sound.stop()
    cap.release()
    cv2.destroyAllWindows()


# ==========================================
# 🎬 影片偵測邏輯
# ==========================================
def detect_video_fn(video_frame, interval, confidence_level, alarm_enabled):
    if video_frame is None:
        return None, generate_html_table([["錯誤", "請先上傳影片", ""]]), None
    
    video_counter = 0
    log_records = []
    cached_boxes = [] 
    is_alerting = False
    sin_token = 0

    get_video = cv2.VideoCapture(video_frame)
    fps = get_video.get(cv2.CAP_PROP_FPS)
    if fps <= 0 or fps is None:
        fps = 24.0

    interval = int(interval) if interval and interval > 0 else 4
    conf_threshold = float(confidence_level) / 100.0 if confidence_level else 0.50

    ret, first_frame = get_video.read()
    if not ret:
        return None, generate_html_table([["錯誤", "無法讀取影片", ""]]), None
        
    h, w = first_frame.shape[:2]
    width = (w // 2) * 2
    height = (h // 2) * 2

    output_dir = "videos"
    os.makedirs(output_dir, exist_ok=True)
    output_name = os.path.join(output_dir, "output_processed.webm")
    output_format = cv2.VideoWriter_fourcc(*"vp80")
    out = cv2.VideoWriter(output_name, output_format, fps, (width, height))
    
    get_video.set(cv2.CAP_PROP_POS_FRAMES, 0)

    while get_video.isOpened():
        end_frame, frame = get_video.read()
        if not end_frame:
            break   

        frame = cv2.resize(frame, (width, height))
        video_counter += 1
        output_frame = frame.copy()

        if (video_counter - 1) % interval == 0:
            results = model(frame, conf=conf_threshold)  
            detected_classes = [model.names[int(cls)].lower() for cls in results[0].boxes.cls]
            
            cached_boxes = []
            for box in results[0].boxes:
                cls_id = int(box.cls[0])
                cls_name = model.names[cls_id]
                conf = float(box.conf[0])
                xyxy = box.xyxy[0].tolist()
                cached_boxes.append((cls_name, conf, xyxy[0], xyxy[1], xyxy[2], xyxy[3]))
            
            if "no_helmet" in detected_classes or "no_gloves" in detected_classes or "no_glove" in detected_classes:
                sin_token += 1
            else:
                if is_alerting:
                    total_seconds = video_counter / fps
                    m = int(total_seconds // 60)
                    s = int(total_seconds % 60)
                    ms = int((total_seconds % 1) * 1000)
                    time_str = f"{m:02d}:{s:02d}.{ms:03d}"
                    log_records.append([time_str, "✅ 裝備已恢復完整，警報解除。", ""])
                    is_alerting = False
                sin_token = 0 
            
            if sin_token == 4:
                is_alerting = True
                total_seconds = video_counter / fps
                m = int(total_seconds // 60)
                s = int(total_seconds % 60)
                ms = int((total_seconds % 1) * 1000)
                time_str = f"{m:02d}:{s:02d}.{ms:03d}"
                
                temp_frame = output_frame.copy()
                draw_cached_boxes(temp_frame, cached_boxes)
                img_name = f"video_alert_{time.strftime('%Y%m%d_%H%M%S')}_{video_counter}.jpg"
                img_path = os.path.join(LOG_IMG_DIR, img_name)
                cv2.imwrite(img_path, temp_frame)
                
                log_records.append([time_str, "⚠️ 偵測到消防員未配戴完整裝備！", img_path])

        draw_cached_boxes(output_frame, cached_boxes)

        if is_alerting and alarm_enabled:
            draw_warning_text(output_frame)

        out.write(output_frame)

    get_video.release()
    out.release()
    
    if not log_records:
        log_records.append(["-", "✅ 影片檢測完畢，未發現違規狀況。", ""])
    
    excel_filename = f"video_log_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    excel_filepath = os.path.join(EXCEL_EXPORT_DIR, excel_filename)
    export_to_excel(log_records, excel_filepath)

    return output_name, generate_html_table(log_records), excel_filepath

# ==========================================
# 網頁 UI 介面
# ==========================================
with gr.Blocks() as demo:
    gr.Markdown('# 消防人員裝備影像偵測系統')

    webcam_history_state = gr.State([])

    with gr.Tab("及時偵測"):
        gr.Markdown("### 📷 本機端攝影機即時偵測")
        start_webcam_btn = gr.Button("🚀 啟動本機攝影機", variant="primary")

    with gr.Tab("影片偵測"):
        video_input = gr.Video(sources=["upload"])
        start_button = gr.Button("開始偵測")
        video_output = gr.Video()

    with gr.Tab("工作日誌"):
        with gr.Tab("及時偵測紀錄"):
            webcam_log_html = gr.HTML(label="檢測紀錄")
            webcam_excel_file = gr.File(label="下載 Excel 日誌檔 (內含圖片)")
            
        with gr.Tab("影像偵測紀錄"):
            video_log_html = gr.HTML(label="檢測紀錄")
            video_excel_file = gr.File(label="下載 Excel 日誌檔 (內含圖片)")

    with gr.Tab("設定"):
        fps_set = gr.Number(label="抽幀間隔 (1 即為每幀皆辨識)", value=1)
        confidence_level = gr.Number(label="信心程度過濾 (0~100)", value=50, interactive=True)
        voiceOff_set = gr.Checkbox(label="啟用警報", value=True)
        dataOff_set = gr.Checkbox(label="啟用資料收集", value=True, interactive=True)
        voiceclass_set = gr.Dropdown(["警報", "吼叫"], label="自訂警報聲", value="警報", interactive=True)
        voice_slider = gr.Slider(minimum=0, maximum=100, value=50, step=1, label="警報聲大小", interactive=True)

    start_webcam_btn.click(
        fn=run_local_webcam,
        inputs=[
            fps_set,
            confidence_level,
            voiceOff_set,
            dataOff_set,
            voiceclass_set,
            voice_slider,
            webcam_history_state
        ],
        outputs=[
            webcam_log_html, 
            webcam_excel_file,
            webcam_history_state 
        ]
    )

    start_button.click(
        fn=detect_video_fn,
        inputs=[
            video_input,
            fps_set,
            confidence_level,
            voiceOff_set,
        ],
        outputs=[
            video_output, 
            video_log_html,
            video_excel_file
        ],
    )

# 🛠️ 修正 6：確保 gradio 允許讀取同層資料夾
demo.launch(allowed_paths=["voice", "data", "log_images"])